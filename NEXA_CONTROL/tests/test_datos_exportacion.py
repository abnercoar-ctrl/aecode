from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from modulos.exportacion import exportar_reporte, libro_excel
from modulos.gestion_datos import (ErrorDatos, cargar_ejemplos, incorporar, leer_archivo, tabla_vacia, validar_datos)
from modulos.indicadores import calcular_indicadores


@pytest.mark.parametrize("campo,valor", [("precio_unitario", -1), ("metrado_contratado", "abc"),
    ("precio_unitario", float("inf")), ("descripcion", " "), ("proyecto", None),
    ("fecha_inicio", "no-fecha"), ("fecha_inicio", 1234), ("fecha_fin", "2025-01-01")])
def test_rechaza_partidas_invalidas(datos, campo, valor):
    p, a = datos
    p[campo] = p[campo].astype(object)
    p.loc[0, campo] = valor
    with pytest.raises(ErrorDatos):
        validar_datos(p, a)


@pytest.mark.parametrize("campo,valor", [("metrado_ejecutado", -1), ("metrado_programado", "abc"),
    ("metrado_ejecutado", 101), ("metrado_programado", 101), ("codigo", "NO-EXISTE"),
    ("proyecto", "OTRO"), ("fecha", "fecha inválida")])
def test_rechaza_avances_invalidos(datos, campo, valor):
    p, a = datos
    a[campo] = a[campo].astype(object)
    a.loc[0, campo] = valor
    with pytest.raises(ErrorDatos):
        validar_datos(p, a)


def test_duplicados_y_orfanos(datos):
    p, a = datos
    with pytest.raises(ErrorDatos, match="duplicadas"):
        validar_datos(pd.concat([p, p.iloc[:1]]), a)
    with pytest.raises(ErrorDatos, match="duplicado"):
        validar_datos(p, pd.concat([a, a.iloc[:1]]))
    with pytest.raises(ErrorDatos, match="inexistente"):
        incorporar(p, a, p.iloc[:0], "partidas", reemplazar=True)
    assert len(p) == 2 and len(a) == 4


def test_clave_compuesta_no_mezcla_proyectos(datos):
    p, a = datos
    otro = p.iloc[:1].copy()
    otro.proyecto = "OTRO"
    otro.precio_unitario = 100
    avance_otro = a.iloc[:1].copy()
    avance_otro.proyecto = "OTRO"
    p, a = validar_datos(pd.concat([p, otro]), pd.concat([a, avance_otro]))
    ind = calcular_indicadores(p, a, "2026-01-01", "2026-01-31")
    assert ind["presupuesto_contratado"] == 13000
    assert ind["ejecutado_acumulado"] == 2400


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_importacion_preserva_codigo_y_exporta(datos, extension):
    p, a = datos
    contenido = p.to_csv(index=False, date_format="%Y-%m-%d").encode("utf-8-sig") if extension == "csv" else libro_excel({"Partidas": p})
    leida = leer_archivo(contenido, f"archivo.{extension}", "partidas")
    assert leida.codigo.tolist() == ["001", "002"]
    p2, a2 = incorporar(tabla_vacia("partidas"), tabla_vacia("avances"), leida, "partidas")
    p2, a2 = incorporar(p2, a2, a, "avances")
    assert calcular_indicadores(p2, a2, "2026-01-01", "2026-02-28")["ejecutado_acumulado"] == 1500


def test_archivos_invalidos(datos):
    with pytest.raises(ErrorDatos, match="Columnas"):
        leer_archivo(b"columna,otra\n1,2", "malo.csv", "partidas")
    with pytest.raises(ErrorDatos):
        leer_archivo(b"corrupto", "malo.xlsx", "partidas")
    with pytest.raises(ErrorDatos):
        leer_archivo(b"corrupto", "malo.pdf", "partidas")


def test_exportacion_reproducible_y_texto_literal(datos):
    p, a = datos
    p.loc[0, "descripcion"] = "=1+1"
    contenido = exportar_reporte(p, a, "2026-02-01", "2026-02-28", "UM", "DEMO")
    libro = load_workbook(BytesIO(contenido))
    assert {"Resumen general", "Partidas contratadas", "Registro de avances", "Valorizaciones", "Indicadores"} <= set(libro.sheetnames)
    assert libro["Partidas contratadas"]["C2"].value == "=1+1"
    assert libro["Partidas contratadas"]["C2"].data_type == "s"
    assert libro["Partidas contratadas"].freeze_panes == "A2"
    assert "UM" in libro["Partidas contratadas"]["G2"].number_format
    exportadas = pd.read_excel(BytesIO(contenido), sheet_name="Registro de avances", dtype={"codigo": str})
    assert len(exportadas) == 4  # conserva enero aunque Desde sea febrero
    indicadores = pd.read_excel(BytesIO(contenido), sheet_name="Indicadores").set_index("indicador")
    assert indicadores.loc["ejecutado_acumulado", "valor"] == 1500
    assert indicadores.loc["ejecutado_periodo", "valor"] == 1100


def test_ejemplos_verificables():
    p, a = cargar_ejemplos(Path(__file__).resolve().parents[1] / "datos")
    i = calcular_indicadores(p, a, "2026-01-01", "2026-03-31")
    assert i["presupuesto_contratado"] == 130000
    assert i["ejecutado_acumulado"] == 60000
    assert i["saldo_por_valorizar"] == 70000
    assert i["avance_programado_pct"] == pytest.approx(66 / 130 * 100)
    assert i["avance_real_pct"] == pytest.approx(60 / 130 * 100)
    assert i["trabajos_atrasados"] == 4
